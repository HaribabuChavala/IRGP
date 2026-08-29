import asyncio
import csv
import hashlib
import io
import json
import os
import re
import secrets
import socket
import string
import uuid
from datetime import datetime, timezone
from urllib.parse import urlparse

import httpx
from fastapi import APIRouter, BackgroundTasks, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, EmailStr, Field
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session

try:
    import oracledb
except Exception:  # pragma: no cover - optional dependency at runtime
    oracledb = None

try:
    import teradatasql
except Exception:  # pragma: no cover - optional dependency at runtime
    teradatasql = None

try:
    from impala.dbapi import connect as impala_connect
except Exception:  # pragma: no cover - optional dependency at runtime
    impala_connect = None

import store
from ai_sql_agent import sql_agent_service
from database import SessionLocal, get_db
from models import AuditEvent, DataSource, ExecutionLog, Notification, Organization, QueryHistory, Report, ReportJob, Reminder, Role, User, UserRole
from deps import (
    SecurityContext,
    get_current_user,
    redis_client,
    require_org_admin,
    require_platform_admin,
)
from email_service import send_temp_password_email

router = APIRouter(prefix="/api/v1")

KEYCLOAK_BASE_URL = os.getenv("KEYCLOAK_BASE_URL", "http://keycloak:8080")
KEYCLOAK_REALM = os.getenv("KEYCLOAK_REALM", "report-platform")
KEYCLOAK_ADMIN_USERNAME = os.getenv("KEYCLOAK_ADMIN_USERNAME", "admin")
KEYCLOAK_ADMIN_PASSWORD = os.getenv("KEYCLOAK_ADMIN_PASSWORD", "change-me-admin")
LIVY_SERVICE_URL = os.getenv("LIVY_SERVICE_URL", "http://livy-service:8090/api/livy")

WRITE_PRIVILEGE_KEYWORDS = {
    "ALL",
    "ALTER",
    "CREATE",
    "DELETE",
    "DROP",
    "INSERT",
    "OWNER",
    "TRUNCATE",
    "UPDATE",
    "WRITE",
}


def generate_temp_password(length: int = 16) -> str:
    alphabet = string.ascii_letters + string.digits
    return "".join(secrets.choice(alphabet) for _ in range(length))


def _resolve_log_organization_id(organization_id: str | None, user_id: str | None) -> str | None:
    if organization_id:
        return organization_id
    if not user_id:
        return None
    try:
        with SessionLocal() as db:
            user_record = db.query(User).filter(or_(User.id == user_id, User.keycloak_user_id == user_id)).first()
            if user_record is not None:
                return user_record.organization_id or user_record.tenant_id
    except Exception:
        pass
    return None


def _resolve_db_user_id(user_id: str | None, *, email: str | None = None, username: str | None = None) -> str | None:
    if not user_id and not email and not username:
        return None
    try:
        with SessionLocal() as db:
            candidate = None
            if user_id:
                candidate = db.query(User).filter(or_(User.id == user_id, User.keycloak_user_id == user_id)).first()
            if candidate is None and email:
                candidate = db.query(User).filter(User.email == email).first()
            if candidate is None and username:
                candidate = db.query(User).filter(User.username == username).first()
            if candidate is not None:
                return candidate.id
    except Exception:
        pass
    return user_id or None


def _record_execution_log(
    *,
    organization_id: str | None,
    user_id: str | None,
    job_id: str | None,
    data_source_id: str | None,
    execution_engine: str | None,
    query_id: str | None,
    stage: str,
    source: str,
    status: str,
    message: str,
    error_details: str | None = None,
    context: dict | None = None,
    email: str | None = None,
    username: str | None = None,
) -> None:
    if not organization_id and not user_id and not job_id and not data_source_id:
        return

    resolved_org_id = _resolve_log_organization_id(organization_id, user_id)
    resolved_user_id = _resolve_db_user_id(user_id, email=email, username=username)

    try:
        with SessionLocal() as db:
            db.add(
                ExecutionLog(
                    organization_id=resolved_org_id or organization_id,
                    user_id=resolved_user_id,
                    job_id=job_id,
                    data_source_id=data_source_id,
                    execution_engine=execution_engine,
                    query_id=query_id,
                    stage=stage,
                    source=source,
                    status=status,
                    message=message[:4000],
                    error_details=(error_details or None)[:4000] if error_details else None,
                    context=context or {},
                )
            )
            db.commit()
    except Exception:
        pass


def map_role_to_keycloak_role(role: str) -> str:
    mapping = {
        "ORG_ADMIN": "REPORT_ADMIN",
        "ORG_USER": "REPORT_USER",
        "PLATFORM_ADMIN": "PLATFORM_ADMIN",
    }
    return mapping.get(role, "REPORT_USER")


def map_db_role_name(role: str) -> str:
    mapping = {
        "ORG_ADMIN": "REPORT_ADMIN",
        "ORG_USER": "REPORT_USER",
        "REPORT_ADMIN": "REPORT_ADMIN",
        "REPORT_USER": "REPORT_USER",
        "PLATFORM_ADMIN": "PLATFORM_ADMIN",
    }
    return mapping.get((role or "").upper(), "REPORT_USER")


def ensure_user_role(db: Session, user_record: User, role_name: str) -> None:
    db_role_name = map_db_role_name(role_name)
    role = db.query(Role).filter(Role.name == db_role_name).first()
    if role is None:
        role = Role(name=db_role_name, description=f"Built-in {db_role_name.lower()} role")
        db.add(role)
        db.flush()

    if not db.query(UserRole).filter_by(user_id=user_record.id, role_id=role.id).first():
        db.add(UserRole(user_id=user_record.id, role_id=role.id))


async def get_keycloak_admin_token() -> str:
    async with httpx.AsyncClient(timeout=15) as client:
        resp = await client.post(
            f"{KEYCLOAK_BASE_URL}/realms/master/protocol/openid-connect/token",
            data={
                "username": KEYCLOAK_ADMIN_USERNAME,
                "password": KEYCLOAK_ADMIN_PASSWORD,
                "grant_type": "password",
                "client_id": "admin-cli",
            },
            headers={"Content-Type": "application/x-www-form-urlencoded"},
        )
        if resp.status_code != 200:
            raise HTTPException(status_code=500, detail=f"Keycloak admin auth failed: {resp.text}")
        return resp.json()["access_token"]


async def create_keycloak_user(
    *,
    username: str,
    email: str,
    full_name: str,
    password: str,
    role: str,
    organization_id: str,
    organization_name: str,
    region: str,
) -> str:
    admin_token = await get_keycloak_admin_token()
    realm_role_name = map_role_to_keycloak_role(role)

    user_payload = {
        "username": username,
        "email": email.lower(),
        "firstName": full_name.split(" ")[0] if full_name else username,
        "lastName": " ".join(full_name.split(" ")[1:]) if full_name and " " in full_name else "",
        "enabled": True,
        "emailVerified": False,
        "attributes": {
            "tenant_id": [organization_id],
            "organization_name": [organization_name],
            "region": [region],
        },
        "credentials": [
            {"type": "password", "value": password, "temporary": True},
        ],
    }

    async with httpx.AsyncClient(timeout=15) as client:
        create_resp = await client.post(
            f"{KEYCLOAK_BASE_URL}/admin/realms/{KEYCLOAK_REALM}/users",
            json=user_payload,
            headers={
                "Authorization": f"Bearer {admin_token}",
                "Content-Type": "application/json",
            },
        )

        if create_resp.status_code not in (201, 409):
            raise HTTPException(status_code=500, detail=f"Failed to create Keycloak user: {create_resp.text}")

        if create_resp.status_code == 409:
            user_list_resp = await client.get(
                f"{KEYCLOAK_BASE_URL}/admin/realms/{KEYCLOAK_REALM}/users?email={email.lower()}",
                headers={"Authorization": f"Bearer {admin_token}"},
            )
            if user_list_resp.status_code != 200 or not user_list_resp.json():
                raise HTTPException(status_code=500, detail=f"Keycloak user conflict for {email}, but lookup failed")
            keycloak_user_id = user_list_resp.json()[0]["id"]
        else:
            keycloak_user_id = create_resp.headers.get("Location", "").split("/")[-1]

        role_resp = await client.get(
            f"{KEYCLOAK_BASE_URL}/admin/realms/{KEYCLOAK_REALM}/roles/{realm_role_name}",
            headers={"Authorization": f"Bearer {admin_token}"},
        )
        if role_resp.status_code != 200:
            raise HTTPException(status_code=500, detail=f"Role lookup failed: {role_resp.text}")

        role_payload = [{
            "id": role_resp.json()["id"],
            "name": role_resp.json()["name"],
            "composite": False,
            "clientRole": False,
            "containerId": KEYCLOAK_REALM,
        }]

        assign_resp = await client.post(
            f"{KEYCLOAK_BASE_URL}/admin/realms/{KEYCLOAK_REALM}/users/{keycloak_user_id}/role-mappings/realm",
            json=role_payload,
            headers={
                "Authorization": f"Bearer {admin_token}",
                "Content-Type": "application/json",
            },
        )
        if assign_resp.status_code not in (200, 204):
            raise HTTPException(status_code=500, detail=f"Role assignment failed: {assign_resp.text}")

        return keycloak_user_id


class OrganizationCreate(BaseModel):
    name: str
    contactEmail: str
    region: str
    plan: str = "free"


class OrganizationUserInvite(BaseModel):
    email: EmailStr
    username: str = Field(min_length=3, max_length=150)
    full_name: str | None = None
    role: str = "ORG_USER"


class OrganizationOnboardRequest(BaseModel):
    organization_name: str = Field(min_length=2, max_length=255)
    contact_email: EmailStr
    region: str = Field(min_length=2, max_length=100)
    plan: str = "starter"
    organization_admin: OrganizationUserInvite
    users: list[OrganizationUserInvite] = Field(default_factory=list)


class OrganizationMemberResult(BaseModel):
    id: str
    email: str
    username: str
    full_name: str | None = None
    role: str
    force_password_reset: bool
    temporary_password: str
    delivery_status: str = "queued"
    delivery_provider: str = "mock"


class OrganizationOnboardingResponse(BaseModel):
    organization: dict
    created_users: list[OrganizationMemberResult]
    email_delivery: str = "queued"


class UserRegister(BaseModel):
    name: str
    email: str
    organizationId: str
    role: str = "ORG_USER"


class DataSourceCreate(BaseModel):
    name: str
    type: str
    testConnectionId: str
    connectionUrl: str | None = None
    host: str | None = None
    port: int | None = None
    database: str | None = None
    schema: str | None = None
    filePath: str | None = None
    username: str | None = None
    password: str | None = None
    accessMode: str = "read"


class DataSourceTestRequest(BaseModel):
    name: str
    type: str
    connectionUrl: str | None = None
    host: str | None = None
    port: int | None = None
    database: str | None = None
    schema: str | None = None
    filePath: str | None = None
    username: str | None = None
    password: str | None = None
    accessMode: str = "read"


def _default_port(source_type: str) -> int | None:
    mapping = {
        "oracle": 1521,
        "teradata": 1025,
        "hive": 10000,
        "jdbc": None,
    }
    return mapping.get((source_type or "").lower())


def _extract_host_port(connection_url: str | None) -> tuple[str | None, int | None]:
    if not connection_url:
        return None, None

    parsed = urlparse(connection_url)
    host = parsed.hostname
    port = parsed.port

    if host:
        return host, port

    # Handle JDBC style values like jdbc:oracle:thin:@//host:1521/service
    jdbc_match = None
    if connection_url.startswith("jdbc:"):
        jdbc_match = urlparse(connection_url.replace("jdbc:", "", 1))
    if jdbc_match and jdbc_match.hostname:
        return jdbc_match.hostname, jdbc_match.port

    return None, None


def _normalize_source_payload(body: DataSourceTestRequest | DataSourceCreate) -> dict:
    source_type = (body.type or "").lower().strip()
    host = (body.host or "").strip() or None
    connection_url = (body.connectionUrl or "").strip() or None

    url_host, url_port = _extract_host_port(connection_url)
    if not host and url_host:
        host = url_host

    port = body.port or url_port or _default_port(source_type)

    return {
        "name": body.name.strip(),
        "type": source_type,
        "connectionUrl": connection_url,
        "host": host,
        "port": port,
        "database": (body.database or "").strip() or None,
        "schema": (body.schema or "").strip() or None,
        "filePath": (body.filePath or "").strip() or None,
        "username": (body.username or "").strip() or None,
        "password": body.password or None,
        "accessMode": (body.accessMode or "").strip().lower() or "unknown",
    }


def _validate_source_payload(payload: dict) -> list[str]:
    errors: list[str] = []
    source_type = payload["type"]

    if not payload["name"]:
        errors.append("Data source name is required")
    if source_type not in {"oracle", "teradata", "hive", "excel", "jdbc"}:
        errors.append(f"Unsupported data source type: {source_type}")

    if payload["accessMode"] != "read":
        errors.append("Only read-only credentials are permitted for report data sources")

    if source_type == "excel":
        if not payload["filePath"]:
            errors.append("filePath is required for Excel sources")
        return errors

    if not payload["connectionUrl"] and not payload["host"]:
        errors.append("Provide either connectionUrl or host")
    if not payload["database"]:
        errors.append("database is required for non-Excel sources")
    if not payload["username"] or not payload["password"]:
        errors.append("Read-only username and password are required")

    return errors


def _credential_risk_warnings(username: str | None) -> list[str]:
    warnings: list[str] = []
    if not username:
        return warnings

    risky = ["admin", "root", "owner", "dba", "write"]
    lower = username.lower()
    if any(token in lower for token in risky):
        warnings.append(
            "Username appears privileged. Use a dedicated read-only account before storing credentials."
        )
    return warnings


def _source_fingerprint(payload: dict) -> str:
    sensitive = {
        "type": payload.get("type"),
        "connectionUrl": payload.get("connectionUrl"),
        "host": payload.get("host"),
        "port": payload.get("port"),
        "database": payload.get("database"),
        "schema": payload.get("schema"),
        "filePath": payload.get("filePath"),
        "username": payload.get("username"),
        "passwordHash": hashlib.sha256((payload.get("password") or "").encode()).hexdigest(),
        "accessMode": payload.get("accessMode"),
    }
    raw = json.dumps(sensitive, sort_keys=True)
    return hashlib.sha256(raw.encode()).hexdigest()


def _ensure_data_source_exists(data_source: dict, tenant_id: str) -> dict:
    if not data_source or not tenant_id or not data_source.get("id"):
        return data_source

    try:
        with SessionLocal() as db:
            row = db.query(DataSource).filter(DataSource.id == data_source["id"]).first()
            if row is None:
                row = DataSource(
                    id=data_source["id"],
                    organization_id=tenant_id,
                    name=data_source.get("name") or "Unnamed data source",
                    type=data_source.get("type") or "unknown",
                    host=data_source.get("host"),
                    port=data_source.get("port"),
                    database_name=data_source.get("database") or data_source.get("databaseName"),
                    file_path=data_source.get("filePath"),
                    status=data_source.get("status") or "connected",
                    query_count=int(data_source.get("queryCount") or 0),
                )
                db.add(row)
            else:
                row.organization_id = tenant_id
                row.name = data_source.get("name") or row.name or "Unnamed data source"
                row.type = data_source.get("type") or row.type or "unknown"
                row.host = data_source.get("host") or row.host
                row.port = data_source.get("port") or row.port
                row.database_name = data_source.get("database") or data_source.get("databaseName") or row.database_name
                row.file_path = data_source.get("filePath") if data_source.get("filePath") is not None else row.file_path
                row.status = data_source.get("status") or row.status or "connected"
                row.query_count = int(data_source.get("queryCount") or row.query_count or 0)

            db.commit()
    except Exception:
        pass

    return data_source


async def _resolve_data_source_credentials(data_source: dict, tenant_id: str) -> dict:
    resolved = dict(data_source)
    if resolved.get("username") and resolved.get("password"):
        return resolved

    vault_path = resolved.get("vaultSecretPath")
    if not vault_path:
        vault_path = f"secret/data/report-platform/{tenant_id}/datasources/{resolved.get('id') or 'unknown'}"

    vault_addr = os.getenv("VAULT_ADDR", "http://vault:8200").rstrip("/")
    vault_token = os.getenv("VAULT_TOKEN", "dev-root-token").strip()

    try:
        async with httpx.AsyncClient(timeout=8) as client:
            response = await client.get(
                f"{vault_addr}/v1/{vault_path.lstrip('/')}",
                headers={"X-Vault-Token": vault_token},
            )
            if response.status_code == 200:
                secret_data = response.json().get("data", {}).get("data", {})
                if secret_data:
                    resolved.update(secret_data)
                    return resolved
    except Exception:
        pass

    if str(resolved.get("type", "")).lower() == "oracle" and (
        str(resolved.get("host") or "").lower() in {"oracle", "localhost"}
        or str(resolved.get("database") or "").upper() == "FREEPDB1"
    ):
        resolved["username"] = "report_ro"
        resolved["password"] = "ReportReadOnly123"
        resolved["schema"] = resolved.get("schema") or "report_owner"
        resolved["database"] = resolved.get("database") or "FREEPDB1"
        resolved["accessMode"] = resolved.get("accessMode") or "read"

    if str(resolved.get("type", "")).lower() == "teradata" and (
        str(resolved.get("host") or "").lower() in {"teradata-lab", "localhost", "127.0.0.1"}
        or str(resolved.get("database") or "").lower() == "teradata_lab"
    ):
        resolved["username"] = "td_readonly"
        resolved["password"] = "TdReadOnly123"
        resolved["database"] = resolved.get("database") or "teradata_lab"
        resolved["schema"] = resolved.get("schema") or "public"
        resolved["accessMode"] = resolved.get("accessMode") or "read"

    return resolved


def _data_source_name_exists(tenant_id: str, candidate_name: str) -> bool:
    normalized = (candidate_name or "").strip().lower()
    if not normalized:
        return False

    sources = store.get_data_sources(redis_client, tenant_id)
    return any((item.get("name") or "").strip().lower() == normalized for item in sources)


async def _tcp_check(host: str, port: int) -> bool:
    def _connect() -> bool:
        with socket.create_connection((host, port), timeout=3):
            return True

    try:
        return await asyncio.to_thread(_connect)
    except OSError:
        return False


def _db_connect_timeout_seconds() -> int:
    raw = os.getenv("DB_TEST_CONNECT_TIMEOUT_SECONDS", "5").strip()
    try:
        value = int(raw)
    except ValueError:
        value = 5
    return max(2, min(value, 30))


def _friendly_connection_error(engine: str, exc: Exception) -> str:
    message = " ".join(str(exc).split())
    lower = message.lower()

    if any(token in lower for token in ["ora-01017", "invalid username/password", "authentication failed"]):
        return f"{engine} login failed. Check the username and password."

    if any(token in lower for token in ["no listener", "connection refused", "cannot reach", "timed out", "timeout", "name or service not known"]):
        return f"{engine} server is not reachable. Check the host, port, and network access."

    if any(token in lower for token in ["service name", "sid", "unknown database", "database does not exist"]):
        return f"{engine} database or service name is invalid. Check the database/service field and connection URL."

    if any(token in lower for token in ["dpy-4011", "dpy-6005", "tsocket read 0 bytes", "lost connection", "failure receiving message header"]):
        return f"{engine} connection failed during handshake. Check that the selected datasource type matches the target server and that the port is correct."

    return f"{engine} connection failed. Verify the host, port, database/service, and read-only credentials."


def _friendly_vault_error(detail: str) -> str:
    return "Credentials could not be saved to Vault. Verify that Vault is running and the API has a valid Vault token."


def _contains_write_privilege(text: str) -> bool:
    upper = (text or "").upper()
    return any(keyword in upper for keyword in WRITE_PRIVILEGE_KEYWORDS)


def _extract_oracle_service_name(payload: dict) -> str | None:
    if payload.get("database"):
        return str(payload["database"])

    connection_url = payload.get("connectionUrl") or ""
    if not connection_url:
        return None

    normalized = connection_url
    if normalized.startswith("jdbc:"):
        normalized = normalized.replace("jdbc:", "", 1)

    parsed = urlparse(normalized)
    if parsed.path:
        service = parsed.path.strip("/")
        if service:
            return service

    marker = "@//"
    if marker in connection_url:
        tail = connection_url.split(marker, 1)[1]
        if "/" in tail:
            return tail.split("/", 1)[1].split("?", 1)[0].strip() or None

    return None


def _oracle_authz_check(payload: dict) -> tuple[bool, str]:
    if oracledb is None:
        return False, "Oracle driver is not installed (oracledb)"

    host = payload.get("host")
    port = int(payload.get("port") or 1521)
    username = payload.get("username")
    password = payload.get("password")
    service_name = _extract_oracle_service_name(payload)

    if not service_name:
        return False, "Oracle service name is required (use database field or include it in connectionUrl)"

    connection = None
    try:
        connection = oracledb.connect(
            user=username,
            password=password,
            host=host,
            port=port,
            service_name=service_name,
            tcp_connect_timeout=_db_connect_timeout_seconds(),
        )
        with connection.cursor() as cursor:
            cursor.execute("SELECT 1 FROM dual")
            cursor.fetchone()

            cursor.execute(
                """
                SELECT COUNT(*)
                FROM SESSION_PRIVS
                WHERE PRIVILEGE IN (
                    'ALTER ANY TABLE',
                    'CREATE ANY TABLE',
                    'DELETE ANY TABLE',
                    'DROP ANY TABLE',
                    'INSERT ANY TABLE',
                    'TRUNCATE ANY TABLE',
                    'UPDATE ANY TABLE'
                )
                """
            )
            write_any_count = int(cursor.fetchone()[0] or 0)
            if write_any_count > 0:
                return False, "Oracle account has broad write privileges; provide read-only credentials"

            cursor.execute(
                """
                SELECT COUNT(*)
                FROM USER_ROLE_PRIVS
                WHERE GRANTED_ROLE IN ('DBA', 'RESOURCE')
                """
            )
            elevated_role_count = int(cursor.fetchone()[0] or 0)
            if elevated_role_count > 0:
                return False, "Oracle account has elevated roles (DBA/RESOURCE); provide read-only credentials"

        return True, "Oracle authentication and read-only privilege checks succeeded"
    except Exception as exc:
        return False, _friendly_connection_error("Oracle", exc)
    finally:
        if connection is not None:
            try:
                connection.close()
            except Exception:
                pass


def _teradata_authz_check(payload: dict) -> tuple[bool, str]:
    test_backend = os.getenv("TERADATA_TEST_BACKEND", "native").strip().lower()
    if test_backend == "postgres":
        try:
            import psycopg
        except Exception as exc:
            return False, f"Teradata lab backend requires psycopg: {str(exc)}"

        host = payload.get("host")
        port = int(payload.get("port") or 5432)
        database = payload.get("database") or "teradata_lab"
        username = payload.get("username")
        password = payload.get("password")

        connection = None
        try:
            connection = psycopg.connect(
                host=host,
                port=port,
                dbname=database,
                user=username,
                password=password,
                connect_timeout=_db_connect_timeout_seconds(),
            )
            with connection.cursor() as cursor:
                cursor.execute("SELECT CURRENT_DATE")
                cursor.fetchone()

                cursor.execute(
                    """
                    SELECT EXISTS (
                        SELECT 1
                        FROM information_schema.table_privileges
                        WHERE grantee = current_user
                          AND privilege_type IN ('INSERT', 'UPDATE', 'DELETE', 'TRUNCATE', 'REFERENCES', 'TRIGGER')
                    )
                    """
                )
                has_write_privileges = bool(cursor.fetchone()[0])
                if has_write_privileges:
                    return False, "Teradata lab account has write privileges; provide read-only credentials"

                cursor.execute("SELECT has_database_privilege(current_user, current_database(), 'CREATE')")
                can_create = bool(cursor.fetchone()[0])
                if can_create:
                    return False, "Teradata lab account can CREATE objects; provide stricter read-only credentials"

            return True, "Teradata lab authentication and read-only privilege checks succeeded"
        except Exception as exc:
            return False, _friendly_connection_error("Teradata lab", exc)
        finally:
            if connection is not None:
                try:
                    connection.close()
                except Exception:
                    pass

    if teradatasql is None:
        return False, "Teradata driver is not installed (teradatasql)"

    params = {
        "host": payload.get("host"),
        "user": payload.get("username"),
        "password": payload.get("password"),
        "logmech": os.getenv("TERADATA_LOGMECH", "TD2"),
        "encryptdata": os.getenv("TERADATA_ENCRYPT_DATA", "true"),
    }
    if payload.get("port"):
        params["dbs_port"] = str(payload["port"])
    if payload.get("database"):
        params["database"] = payload["database"]

    connection = None
    try:
        connection = teradatasql.connect(**params)
        with connection.cursor() as cursor:
            cursor.execute("SELECT CURRENT_DATE")
            cursor.fetchone()

            privilege_queries = [
                """
                SELECT TOP 1 AccessRight
                FROM DBC.AllRightsV
                WHERE Grantee = USER
                  AND AccessRight IN ('CD', 'CT', 'D', 'DL', 'DM', 'DR', 'DT', 'I', 'U')
                """,
                """
                SELECT TOP 1 AccessRight
                FROM DBC.AllRightsV
                WHERE UserName = USER
                  AND AccessRight IN ('CD', 'CT', 'D', 'DL', 'DM', 'DR', 'DT', 'I', 'U')
                """,
                """
                SELECT TOP 1 AccessRight
                FROM DBC.UserRightsV
                WHERE Grantee = USER
                  AND AccessRight IN ('CD', 'CT', 'D', 'DL', 'DM', 'DR', 'DT', 'I', 'U')
                """,
            ]

            introspection_ran = False
            for query in privilege_queries:
                try:
                    cursor.execute(query)
                    introspection_ran = True
                    row = cursor.fetchone()
                    if row:
                        return False, "Teradata account appears to have write privileges; provide read-only credentials"
                    break
                except Exception:
                    continue

            if not introspection_ran:
                return False, "Teradata privilege introspection failed; grant metadata-read access or use stricter read-only account"

        return True, "Teradata authentication and read-only privilege checks succeeded"
    except Exception as exc:
        return False, _friendly_connection_error("Teradata", exc)
    finally:
        if connection is not None:
            try:
                connection.close()
            except Exception:
                pass


def _hive_authz_check(payload: dict) -> tuple[bool, str]:
    if impala_connect is None:
        return False, "Hive/Impala driver is not installed (impyla)"

    host = payload.get("host")
    port = int(payload.get("port") or 10000)
    username = payload.get("username")
    password = payload.get("password")
    database = payload.get("database") or "default"

    auth_mechanism = os.getenv("HIVE_AUTH_MECHANISM", "NOSASL")

    params = {
        "host": host,
        "port": port,
        "user": username,
        "database": database,
        "auth_mechanism": auth_mechanism,
        "timeout": _db_connect_timeout_seconds(),
    }
    if password:
        params["password"] = password

    connection = None
    try:
        connection = impala_connect(**params)
        with connection.cursor() as cursor:
            cursor.execute("SELECT 1")
            cursor.fetchone()

            safe_user = (username or "").replace("`", "")
            safe_db = database.replace("`", "")
            grant_queries = [
                f"SHOW GRANT USER `{safe_user}` ON DATABASE `{safe_db}`",
                f"SHOW GRANT USER {safe_user} ON DATABASE {safe_db}",
                f"SHOW GRANT USER `{safe_user}` ON SERVER",
                f"SHOW GRANT USER {safe_user} ON SERVER",
            ]

            introspection_ran = False
            for query in grant_queries:
                try:
                    cursor.execute(query)
                    introspection_ran = True
                    rows = cursor.fetchall() or []
                    for row in rows:
                        row_text = " ".join(str(item) for item in row)
                        if _contains_write_privilege(row_text):
                            return False, "Hive account appears to have write privileges; provide read-only credentials"
                    break
                except Exception:
                    continue

            if not introspection_ran:
                return False, "Hive privilege introspection failed; grant SHOW GRANT access or use stricter read-only account"

        return True, "Hive authentication and read-only privilege checks succeeded"
    except Exception as exc:
        return False, _friendly_connection_error("Hive", exc)
    finally:
        if connection is not None:
            try:
                connection.close()
            except Exception:
                pass


async def _test_connection(payload: dict) -> tuple[bool, str]:
    source_type = payload["type"]
    if source_type == "excel":
        return True, "File path format validated for Excel source"

    host = payload.get("host")
    port = payload.get("port")
    if not host or not port:
        return False, "Unable to resolve host/port from connection details"

    ok = await _tcp_check(host, int(port))
    if not ok:
        return False, f"Cannot reach {host}:{port}. Verify network access and connection settings"

    if source_type == "oracle":
        return await asyncio.to_thread(_oracle_authz_check, payload)

    if source_type == "teradata":
        return await asyncio.to_thread(_teradata_authz_check, payload)

    if source_type == "hive":
        return await asyncio.to_thread(_hive_authz_check, payload)

    return True, f"Network connection test succeeded for {host}:{port}"


def _vault_env() -> tuple[str, str]:
    vault_addr = os.getenv("VAULT_ADDR", "http://vault:8200").rstrip("/")
    vault_token = os.getenv("VAULT_TOKEN", "").strip()
    if not vault_token:
        raise HTTPException(status_code=500, detail="Vault token not configured")
    return vault_addr, vault_token


async def _vault_write(path: str, data: dict) -> None:
    vault_addr, vault_token = _vault_env()
    url = f"{vault_addr}/v1/{path.lstrip('/')}"
    payload = {"data": data}

    async with httpx.AsyncClient(timeout=8) as client:
        response = await client.post(
            url,
            json=payload,
            headers={"X-Vault-Token": vault_token, "Content-Type": "application/json"},
        )
    if response.status_code >= 300:
        raise HTTPException(status_code=502, detail=_friendly_vault_error(response.text))


class ReportGenerateRequest(BaseModel):
    prompt: str = Field(min_length=1)
    dataSourceId: str
    executionEngine: str = Field(default="spark")


class SubscriptionUpgradeRequest(BaseModel):
    plan: str


@router.get("/platform/stats")
async def platform_stats(
    user: SecurityContext = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_platform_admin(user)

    total_organizations = db.scalar(select(func.count(Organization.id))) or 0
    active_jobs_rows = db.execute(
        select(
            ReportJob.id,
            Organization.name.label("organization_name"),
            ReportJob.status,
            ReportJob.progress,
            ReportJob.created_at,
        )
        .join(Organization, Organization.id == ReportJob.organization_id)
        .where(ReportJob.status.in_(["queued", "running"]))
        .order_by(ReportJob.created_at.desc())
        .limit(20)
    ).all()

    active_jobs = []
    for row in active_jobs_rows:
        job_id, org_name, status, progress, started_at = row
        active_jobs.append(
            {
                "id": str(job_id),
                "organization": org_name,
                "type": "Report generation",
                "progress": int(progress or 0),
                "status": status,
                "startedAt": started_at.isoformat() if started_at else "just now",
            }
        )

    queries_today = db.scalar(
        select(func.coalesce(func.sum(QueryHistory.row_count), 0))
        .where(QueryHistory.executed_at >= func.date_trunc("day", func.now()))
    ) or 0

    utilization = min(95, 40 + len(active_jobs) * 8 + total_organizations * 2)
    stats = {
        "totalOrganizations": int(total_organizations),
        "activeJobs": len(active_jobs),
        "queriesToday": int(queries_today),
        "utilizationPercent": int(utilization),
    }
    return {"stats": stats, "activeJobs": active_jobs}


@router.get("/platform/organizations")
async def list_organizations(
    user: SecurityContext = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_platform_admin(user)

    rows = db.execute(
        select(
            Organization.id,
            Organization.name,
            Organization.plan,
            Organization.status,
            Organization.region,
            Organization.contact_email,
            Organization.created_at,
            func.count(func.distinct(User.id)).label("user_count"),
            func.count(func.distinct(DataSource.id)).label("data_source_count"),
            func.coalesce(func.sum(QueryHistory.row_count), 0).label("queries_this_month"),
        )
        .outerjoin(User, User.organization_id == Organization.id)
        .outerjoin(DataSource, DataSource.organization_id == Organization.id)
        .outerjoin(QueryHistory, QueryHistory.organization_id == Organization.id)
        .group_by(
            Organization.id,
            Organization.name,
            Organization.plan,
            Organization.status,
            Organization.region,
            Organization.contact_email,
            Organization.created_at,
        )
        .order_by(Organization.name.asc())
    ).all()

    organizations = []
    for row in rows:
        organizations.append(
            {
                "id": str(row[0]),
                "name": row[1],
                "plan": row[2],
                "status": row[3],
                "userCount": int(row[7] or 0),
                "dataSourceCount": int(row[8] or 0),
                "queriesThisMonth": int(row[9] or 0),
                "createdAt": row[6].date().isoformat() if row[6] else None,
                "region": row[4],
                "contactEmail": row[5],
            }
        )

    return {"organizations": organizations}


@router.post("/platform/organizations")
async def register_organization(
    body: OrganizationCreate,
    user: SecurityContext = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_platform_admin(user)

    org = Organization(
        name=body.name,
        plan=body.plan,
        status="active",
        region=body.region,
        contact_email=body.contactEmail,
    )
    db.add(org)
    db.commit()
    db.refresh(org)

    org_payload = {
        "id": str(org.id),
        "name": org.name,
        "plan": org.plan,
        "status": org.status,
        "userCount": 0,
        "dataSourceCount": 0,
        "queriesThisMonth": 0,
        "createdAt": org.created_at.date().isoformat(),
        "region": org.region,
        "contactEmail": org.contact_email,
    }

    orgs = store.get_organizations(redis_client)
    orgs.append(org_payload)
    store.save_organizations(redis_client, orgs)
    store.save_data_sources(redis_client, str(org.id), [])
    return {"organization": org_payload}


@router.post("/platform/organizations/onboard", response_model=OrganizationOnboardingResponse)
async def onboard_organization(
    body: OrganizationOnboardRequest,
    user: SecurityContext = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_platform_admin(user)
    orgs = store.get_organizations(redis_client)
    normalized_name = body.organization_name.strip()
    if any(org.get("name", "").lower() == normalized_name.lower() for org in orgs):
        raise HTTPException(status_code=409, detail="Organization name already exists")

    invited_users = [body.organization_admin, *body.users]
    seen_emails: set[str] = set()
    for invite in invited_users:
        email = invite.email.lower()
        if email in seen_emails:
            raise HTTPException(status_code=409, detail=f"User already invited: {email}")
        seen_emails.add(email)

    org = Organization(
        name=normalized_name,
        plan=body.plan,
        status="active",
        region=body.region,
        contact_email=body.contact_email.lower(),
    )
    db.add(org)
    db.commit()
    db.refresh(org)

    created_users: list[OrganizationMemberResult] = []
    delivery_statuses: list[str] = []
    for invite in invited_users:
        email = invite.email.lower()
        temp_password = "".join(secrets.choice(string.ascii_letters + string.digits) for _ in range(16))
        delivery = send_temp_password_email(
            recipient=email,
            username=invite.username,
            temp_password=temp_password,
            organization_name=org.name,
        )
        delivery_statuses.append(delivery.get("status", "queued"))

        user_record = OrganizationMemberResult(
            id=str(uuid.uuid4()),
            email=email,
            username=invite.username,
            full_name=invite.full_name or invite.username,
            role=invite.role,
            force_password_reset=True,
            temporary_password=temp_password,
            delivery_status=delivery.get("status", "queued"),
            delivery_provider=delivery.get("provider", "mock"),
        )
        created_users.append(user_record)

        db_user = User(
            email=email,
            username=invite.username,
            full_name=invite.full_name or invite.username,
            organization_id=str(org.id),
            tenant_id=str(org.id),
            organization_name=org.name,
            region=body.region,
            status="invited",
        )
        db.add(db_user)
        db.flush()
        ensure_user_role(db, db_user, invite.role)

    db.commit()

    org_payload = {
        "id": str(org.id),
        "name": org.name,
        "plan": org.plan,
        "status": org.status,
        "userCount": len(created_users),
        "dataSourceCount": 0,
        "queriesThisMonth": 0,
        "createdAt": org.created_at.date().isoformat(),
        "region": org.region,
        "contactEmail": org.contact_email,
    }
    orgs.append(org_payload)
    store.save_organizations(redis_client, orgs)
    store.save_data_sources(redis_client, str(org.id), [])

    return {
        "organization": org_payload,
        "created_users": created_users,
        "email_delivery": "sent" if any(status == "sent" for status in delivery_statuses) else "queued",
    }


@router.post("/platform/users")
async def register_user(
    body: UserRegister,
    user: SecurityContext = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    require_platform_admin(user)

    orgs = store.get_organizations(redis_client)
    org = next((o for o in orgs if o["id"] == body.organizationId), None)
    if not org:
        raise HTTPException(status_code=404, detail="Organization not found")

    email = body.email.lower()
    username = body.email.split("@")[0].lower().replace(" ", "_") if "@" in body.email else body.email.lower()
    temp_password = generate_temp_password()
    delivery = send_temp_password_email(
        recipient=email,
        username=username,
        temp_password=temp_password,
        organization_name=org["name"],
    )

    keycloak_user_id = await create_keycloak_user(
        username=username,
        email=email,
        full_name=body.name,
        password=temp_password,
        role=body.role,
        organization_id=body.organizationId,
        organization_name=org["name"],
        region=org.get("region", "unknown"),
    )

    db_user = User(
        email=email,
        username=username,
        full_name=body.name,
        organization_id=body.organizationId,
        tenant_id=body.organizationId,
        organization_name=org["name"],
        region=org.get("region", "unknown"),
        status="invited",
        keycloak_user_id=keycloak_user_id,
    )
    db.add(db_user)
    db.flush()
    ensure_user_role(db, db_user, body.role)
    db.commit()
    db.refresh(db_user)

    org["userCount"] = org.get("userCount", 0) + 1
    store.save_organizations(redis_client, orgs)

    return {
        "user": {
            "id": str(db_user.id),
            "name": body.name,
            "username": username,
            "email": email,
            "organizationId": body.organizationId,
            "role": body.role,
            "status": "invited",
            "force_password_reset": True,
            "temporary_password": temp_password,
            "delivery_status": delivery.get("status", "queued"),
            "delivery_provider": delivery.get("provider", "mock"),
        },
        "note": "User created and temp password email sent",
    }


@router.get("/org/dashboard")
async def org_dashboard(user: SecurityContext = Depends(get_current_user)):
    tenant_id = user.tenant_id
    if not tenant_id:
        raise HTTPException(status_code=400, detail="tenant_id missing from token")

    sources = [
        _ensure_data_source_exists(source, tenant_id)
        for source in store.get_data_sources(redis_client, tenant_id)
    ]
    history = store.get_query_history(redis_client, tenant_id)
    total_queries = sum(s.get("queryCount", 0) for s in sources)

    return {
        "totalQueries": total_queries,
        "dataSourceCount": len(sources),
        "recentExecutions": len(history),
        "dataSources": sources,
        "queryHistory": history,
    }


@router.get("/org/logs")
async def get_execution_logs(
    user: SecurityContext = Depends(get_current_user),
    db: Session = Depends(get_db),
    limit: int = 200,
):
    tenant_id = user.tenant_id
    if not tenant_id:
        raise HTTPException(status_code=400, detail="tenant_id missing from token")

    current_user = db.query(User).filter(or_(User.email == user.email, User.username == user.username, User.id == user.user_id, User.keycloak_user_id == user.user_id)).first()
    user_key_match = current_user.id if current_user else None

    logs = (
        db.query(ExecutionLog)
        .filter(
            or_(
                ExecutionLog.organization_id == tenant_id,
                ExecutionLog.user_id == user.user_id,
                ExecutionLog.user_id == user_key_match,
            )
        )
        .order_by(ExecutionLog.created_at.desc())
        .limit(limit)
        .all()
    )

    return {
        "logs": [
            {
                "id": item.id,
                "jobId": item.job_id,
                "dataSourceId": item.data_source_id,
                "executionEngine": item.execution_engine,
                "queryId": item.query_id,
                "stage": item.stage,
                "source": item.source,
                "status": item.status,
                "message": item.message,
                "errorDetails": item.error_details,
                "context": item.context,
                "createdAt": item.created_at.isoformat(),
            }
            for item in logs
        ]
    }


@router.get("/org/data-sources")
async def list_data_sources(user: SecurityContext = Depends(get_current_user)):
    tenant_id = user.tenant_id
    if not tenant_id:
        raise HTTPException(status_code=400, detail="tenant_id missing from token")
    sources = [
        _ensure_data_source_exists(source, tenant_id)
        for source in store.get_data_sources(redis_client, tenant_id)
    ]
    return {"dataSources": sources}


@router.post("/org/data-sources/test-connection")
async def test_data_source_connection(
    body: DataSourceTestRequest,
    user: SecurityContext = Depends(get_current_user),
):
    require_org_admin(user)
    tenant_id = user.tenant_id
    if not tenant_id:
        raise HTTPException(status_code=400, detail="tenant_id missing from token")

    payload = _normalize_source_payload(body)
    if _data_source_name_exists(tenant_id, payload.get("name") or ""):
        raise HTTPException(status_code=409, detail="Connection name already exists in this organization. Use a unique name")

    errors = _validate_source_payload(payload)
    if errors:
        raise HTTPException(status_code=422, detail="; ".join(errors))

    warnings = _credential_risk_warnings(payload.get("username"))
    if warnings:
        raise HTTPException(status_code=422, detail="; ".join(warnings))

    success, message = await _test_connection(payload)
    if not success:
        raise HTTPException(status_code=400, detail=message)

    test_id = str(uuid.uuid4())
    redis_key = f"tenant:{tenant_id}:data_source_test:{test_id}"
    redis_client.set(
        redis_key,
        json.dumps({"fingerprint": _source_fingerprint(payload), "createdAt": datetime.now(timezone.utc).isoformat()}),
        ex=600,
    )

    return {
        "success": True,
        "testConnectionId": test_id,
        "normalized": {
            "host": payload.get("host"),
            "port": payload.get("port"),
            "connectionUrl": payload.get("connectionUrl"),
            "database": payload.get("database"),
            "schema": payload.get("schema"),
        },
        "warnings": [],
        "message": message,
    }


@router.post("/org/data-sources")
async def create_data_source(
    body: DataSourceCreate,
    user: SecurityContext = Depends(get_current_user),
):
    require_org_admin(user)
    tenant_id = user.tenant_id
    if not tenant_id:
        raise HTTPException(status_code=400, detail="tenant_id missing from token")

    payload = _normalize_source_payload(body)
    if _data_source_name_exists(tenant_id, payload.get("name") or ""):
        raise HTTPException(status_code=409, detail="Connection name already exists in this organization. Use a unique name")

    errors = _validate_source_payload(payload)
    if errors:
        raise HTTPException(status_code=422, detail="; ".join(errors))

    warnings = _credential_risk_warnings(payload.get("username"))
    if warnings:
        raise HTTPException(status_code=422, detail="; ".join(warnings))

    test_key = f"tenant:{tenant_id}:data_source_test:{body.testConnectionId}"
    saved_test = redis_client.get(test_key)
    if not saved_test:
        raise HTTPException(status_code=400, detail="Connection test expired or missing. Test connection again before saving")

    test_payload = json.loads(saved_test)
    if test_payload.get("fingerprint") != _source_fingerprint(payload):
        raise HTTPException(status_code=400, detail="Connection details changed after test. Re-run test connection")

    source_id = str(uuid.uuid4())
    credential_path = f"secret/data/report-platform/{tenant_id}/datasources/{source_id}"
    await _vault_write(
        credential_path,
        {
            "username": payload.get("username"),
            "password": payload.get("password"),
            "access_mode": payload.get("accessMode"),
            "source_type": payload.get("type"),
            "created_by": user.user_id,
        },
    )

    sources = store.get_data_sources(redis_client, tenant_id)
    source = {
        "id": source_id,
        "name": payload.get("name"),
        "type": payload.get("type"),
        "connectionUrl": payload.get("connectionUrl"),
        "host": payload.get("host"),
        "port": payload.get("port"),
        "database": payload.get("database"),
        "schema": payload.get("schema"),
        "filePath": payload.get("filePath"),
        "accessMode": payload.get("accessMode"),
        "vaultSecretPath": credential_path,
        "status": "connected",
        "lastValidatedAt": datetime.now(timezone.utc).isoformat(),
        "queryCount": 0,
    }
    sources.append(source)
    store.save_data_sources(redis_client, tenant_id, sources)
    _ensure_data_source_exists(source, tenant_id)
    redis_client.delete(test_key)

    orgs = store.get_organizations(redis_client)
    for org in orgs:
        if org["id"] == tenant_id:
            org["dataSourceCount"] = len(sources)
            break
    store.save_organizations(redis_client, orgs)

    return {"dataSource": source}


@router.get("/org/notifications")
async def list_notifications(user: SecurityContext = Depends(get_current_user)):
    tenant_id = user.tenant_id
    if not tenant_id:
        raise HTTPException(status_code=400, detail="tenant_id missing from token")

    redis_notifications = store.get_notifications(redis_client, tenant_id)
    with SessionLocal() as db:
        db_notifications = [
            {
                "id": row.id,
                "title": row.title,
                "message": row.message,
                "type": row.type,
                "read": row.is_read,
                "createdAt": row.created_at.isoformat(),
            }
            for row in db.query(Notification)
            .filter(Notification.organization_id == tenant_id)
            .order_by(Notification.created_at.desc())
            .limit(50)
            .all()
        ]
    if db_notifications:
        return {"notifications": db_notifications}
    return {"notifications": redis_notifications}


@router.patch("/org/notifications/{notification_id}/read")
async def mark_notification_read(
    notification_id: str,
    user: SecurityContext = Depends(get_current_user),
):
    tenant_id = user.tenant_id
    items = store.get_notifications(redis_client, tenant_id)
    for item in items:
        if item["id"] == notification_id:
            item["read"] = True
            break
    store.save_notifications(redis_client, tenant_id, items)

    with SessionLocal() as db:
        record = db.query(Notification).filter(Notification.id == notification_id).first()
        if record is not None:
            record.is_read = True
            db.commit()
    return {"updated": True}


@router.get("/org/reminders")
async def list_reminders(user: SecurityContext = Depends(get_current_user)):
    tenant_id = user.tenant_id
    if not tenant_id:
        raise HTTPException(status_code=400, detail="tenant_id missing from token")
    redis_reminders = store.get_subscription_reminders(redis_client, tenant_id)

    with SessionLocal() as db:
        db_reminders = [
            {
                "id": row.id,
                "title": row.title,
                "message": row.message,
                "type": row.reminder_type,
                "dueAt": row.due_at.isoformat() if row.due_at else None,
                "read": row.is_sent,
                "createdAt": row.created_at.isoformat(),
            }
            for row in db.query(Reminder)
            .filter(Reminder.organization_id == tenant_id)
            .order_by(Reminder.created_at.desc())
            .limit(50)
            .all()
        ]

    if db_reminders:
        return {"reminders": db_reminders}
    return {"reminders": redis_reminders}


@router.get("/org/subscriptions/plans")
async def subscription_plans(user: SecurityContext = Depends(get_current_user)):
    return {"plans": store.SUBSCRIPTION_PLANS}


@router.get("/org/subscriptions/current")
async def current_subscription(user: SecurityContext = Depends(get_current_user)):
    org = store.get_org_by_tenant(redis_client, user.tenant_id)
    if not org:
        return {"plan": "free", "organizationName": user.organization_name}
    return {"plan": org.get("plan", "free"), "organization": org}


@router.post("/org/subscriptions/upgrade")
async def upgrade_subscription(
    body: SubscriptionUpgradeRequest,
    user: SecurityContext = Depends(get_current_user),
):
    require_org_admin(user)
    orgs = store.get_organizations(redis_client)
    updated = None
    for org in orgs:
        if org["id"] == user.tenant_id:
            org["plan"] = body.plan
            updated = org
            break
    if not updated:
        raise HTTPException(status_code=404, detail="Organization not found")
    store.save_organizations(redis_client, orgs)
    return {"organization": updated}


def _sql_result_to_table(rows: list[tuple], columns: list[str] | None = None) -> list[list[str]]:
    if not rows:
        return [columns or ["result"], ["No rows returned"]]

    header = [str(column[0] if isinstance(column, tuple) else column) for column in (columns or [])]
    if not header and hasattr(rows[0], "_fields"):
        header = [str(field) for field in rows[0]._fields]
    if not header:
        header = [f"Column {index + 1}" for index in range(len(rows[0]))]

    result = [header]
    for row in rows:
        result.append(["" if value is None else str(value) for value in row])
    return result


def _normalize_execution_engine(value: str | None, source_type: str | None = None) -> str:
    engine = (value or "").strip().lower()
    if engine in {"spark", "oracle", "teradata"}:
        return engine

    source = (source_type or "").strip().lower()
    if source in {"oracle", "teradata"}:
        return source

    return "spark"


def _execute_report_sql(sql: str, data_source: dict) -> list[list[str]]:
    source_type = (data_source.get("type") or "").lower()
    host = data_source.get("host") or "localhost"
    port = int(data_source.get("port") or 1521)
    username = data_source.get("username") or ""
    password = data_source.get("password") or ""
    database = data_source.get("database") or data_source.get("schema") or ""

    if source_type == "oracle":
        if oracledb is None:
            raise RuntimeError("Oracle driver is not installed")
        service_name = database or "FREEPDB1"
        connection = oracledb.connect(
            user=username,
            password=password,
            host=host,
            port=port,
            service_name=service_name,
            tcp_connect_timeout=_db_connect_timeout_seconds(),
        )
        try:
            cursor = connection.cursor()
            try:
                cursor.execute(sql)
                rows = cursor.fetchall()
                columns = [column[0] for column in cursor.description] if cursor.description else None
                return _sql_result_to_table(rows, columns)
            finally:
                cursor.close()
        finally:
            connection.close()

    if source_type == "teradata":
        try:
            import psycopg
        except Exception as exc:  # pragma: no cover
            raise RuntimeError(f"Teradata lab driver is not installed: {exc}") from exc
        connection = psycopg.connect(
            host=host,
            port=port,
            dbname=database or "teradata_lab",
            user=username,
            password=password,
            connect_timeout=10,
        )
        try:
            with connection.cursor() as cursor:
                cursor.execute(sql)
                rows = cursor.fetchall()
                columns = [desc[0] for desc in cursor.description] if cursor.description else None
                return _sql_result_to_table(rows, columns)
        finally:
            connection.close()

    raise RuntimeError(f"Datasource type '{data_source.get('type')}' is not supported for query execution")


def _spark_safe_sql(sql: str) -> str:
    if not isinstance(sql, str):
        return sql

    cleaned = sql.strip()
    if not cleaned:
        return cleaned

    # Spark does not understand Oracle's DUAL pseudo-table, so replace it with a tiny inline table.
    if re.search(r"\bFROM\s+dual\b", cleaned, flags=re.IGNORECASE):
        cleaned = re.sub(r"\bFROM\s+dual\b", "FROM (SELECT 1 AS _dual_dummy) AS dual", cleaned, flags=re.IGNORECASE)

    # Oracle SQL and ANSI dialects commonly use FETCH FIRST n ROWS ONLY.
    # Spark expects LIMIT n, which is the same semantic when used after ORDER BY.
    cleaned = re.sub(
        r"\bFETCH\s+FIRST\s+(\d+)\s+ROWS\s+ONLY\b",
        r"LIMIT \1",
        cleaned,
        flags=re.IGNORECASE,
    )

    return cleaned


def _extract_livy_session_id(session_payload: object) -> int | None:
    if not isinstance(session_payload, dict):
        return None

    for key in ("id", "sessionId", "session_id"):
        value = session_payload.get(key)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            return int(value)
        if isinstance(value, str) and value.strip():
            try:
                return int(value)
            except ValueError:
                pass

    for key in ("session", "result", "data"):
        nested = session_payload.get(key)
        if isinstance(nested, dict):
            nested_id = _extract_livy_session_id(nested)
            if nested_id is not None:
                return nested_id

    return None


async def _execute_report_sql_via_livy(
    sql: str,
    data_source: dict,
    *,
    organization_id: str | None = None,
    user_id: str | None = None,
    job_id: str | None = None,
    data_source_id: str | None = None,
    execution_engine: str = "spark",
) -> list[list[str]]:
    source_type = (data_source.get("type") or "").lower()
    connection_url = (data_source.get("connectionUrl") or "").strip()
    host = data_source.get("host") or "localhost"
    port = data_source.get("port") or 1521
    username = data_source.get("username") or ""
    password = data_source.get("password") or ""
    database = data_source.get("database") or data_source.get("schema") or "FREEPDB1"
    sql = _spark_safe_sql(sql)

    connector_type = "oracle" if source_type == "oracle" else "hive"
    payload = {
        "kind": "spark",
        "connectorType": connector_type,
        "jdbcUrl": connection_url or f"jdbc:oracle:thin:@//{host}:{port}/{database}",
        "user": username,
        "password": password,
        "partitionColumn": "1",
        "numPartitions": 1,
        "sql": sql,
    }
    _record_execution_log(
        organization_id=organization_id,
        user_id=user_id,
        job_id=job_id,
        data_source_id=data_source_id,
        execution_engine=execution_engine,
        query_id=f"{job_id or 'spark'}-query",
        stage="livy_request",
        source="backend",
        status="info",
        message="Sending Spark SQL to Livy session",
        context={"connectorType": connector_type, "jdbcUrl": connection_url or f"jdbc:oracle:thin:@//{host}:{port}/{database}", "sql": sql},
    )

    def _extract_rows_from_output(output: dict) -> list[list[str]]:
        if not isinstance(output, dict):
            return []

        data_section = output.get("data") or {}
        if isinstance(data_section, dict):
            json_value = data_section.get("application/json")
            if isinstance(json_value, str):
                try:
                    parsed = json.loads(json_value)
                except json.JSONDecodeError:
                    parsed = None
                if isinstance(parsed, list) and parsed:
                    if isinstance(parsed[0], dict):
                        headers = [str(k) for k in parsed[0].keys()]
                        rows = [[str(item.get(h, "")) for h in headers] for item in parsed]
                        return [headers, *rows]
                    if isinstance(parsed[0], list):
                        return [list(map(str, parsed[0]))] + [list(map(str, row)) for row in parsed[1:]]

            text_value = data_section.get("text/plain")
            if isinstance(text_value, str) and text_value.strip():
                output_for_parse = {"data": {"text/plain": text_value}}
                parsed_rows = _extract_rows_from_output(output_for_parse)
                if parsed_rows:
                    return parsed_rows

        text_value = output.get("text")
        if not isinstance(text_value, str) and isinstance(output.get("data"), dict):
            text_value = output["data"].get("text/plain")
        if isinstance(text_value, str) and text_value.strip():
            lines = [line.strip() for line in text_value.splitlines() if line.strip()]
            table_rows: list[list[str]] = []
            for line in lines:
                if "|" not in line or line.startswith("+") or line.startswith("-"):
                    continue
                row_values = [cell.strip() for cell in line.strip("|").split("|")]
                row_values = [cell for cell in row_values if cell != ""]
                if not row_values or len(row_values) == 1 and row_values[0].startswith(("spark:", "df:", "res")):
                    continue
                table_rows.append(row_values)
            if len(table_rows) >= 2:
                return table_rows

        return []

    async with httpx.AsyncClient(timeout=60) as client:
        response = await client.post(f"{LIVY_SERVICE_URL}/sessions", json=payload)
        if response.status_code >= 300:
            error_message = f"Livy session creation failed ({response.status_code}): {response.text[:500]}"
            _record_execution_log(
                organization_id=organization_id,
                user_id=user_id,
                job_id=job_id,
                data_source_id=data_source_id,
                execution_engine=execution_engine,
                query_id=f"{job_id or 'spark'}-query",
                stage="livy_session",
                source="backend",
                status="error",
                message="Livy session creation failed",
                error_details=error_message,
                context={"payload": payload},
            )
            raise RuntimeError(error_message)

        session = response.json()
        session_id = _extract_livy_session_id(session)
        if session_id is None:
            error_message = f"Livy session response did not include an ID: {session}"
            _record_execution_log(
                organization_id=organization_id,
                user_id=user_id,
                job_id=job_id,
                data_source_id=data_source_id,
                execution_engine=execution_engine,
                query_id=f"{job_id or 'spark'}-query",
                stage="livy_session",
                source="backend",
                status="error",
                message="Livy session response missing session ID",
                error_details=error_message,
                context={"response": session},
            )
            raise RuntimeError(error_message)

        _record_execution_log(
            organization_id=organization_id,
            user_id=user_id,
            job_id=job_id,
            data_source_id=data_source_id,
            execution_engine=execution_engine,
            query_id=f"{job_id or 'spark'}-query",
            stage="livy_session",
            source="backend",
            status="success",
            message="Livy session created successfully",
            context={"sessionId": session_id, "state": session.get("state")},
        )

        session_state = (session.get("state") or "").lower()
        while session_state not in {"idle", "error", "dead"}:
            await asyncio.sleep(2)
            status_response = await client.get(f"{LIVY_SERVICE_URL}/sessions/{session_id}")
            if status_response.status_code >= 300:
                error_message = f"Livy session status check failed ({status_response.status_code}): {status_response.text[:500]}"
                _record_execution_log(
                    organization_id=organization_id,
                    user_id=user_id,
                    job_id=job_id,
                    data_source_id=data_source_id,
                    execution_engine=execution_engine,
                    query_id=f"{job_id or 'spark'}-query",
                    stage="livy_session",
                    source="backend",
                    status="error",
                    message="Livy session status check failed",
                    error_details=error_message,
                    context={"sessionId": session_id},
                )
                raise RuntimeError(error_message)
            session = status_response.json()
            session_state = (session.get("state") or "").lower()
        if session_state in {"error", "dead"}:
            error_message = f"Livy session ended in an unusable state: {session_state}"
            _record_execution_log(
                organization_id=organization_id,
                user_id=user_id,
                job_id=job_id,
                data_source_id=data_source_id,
                execution_engine=execution_engine,
                query_id=f"{job_id or 'spark'}-query",
                stage="livy_session",
                source="backend",
                status="error",
                message="Livy session unusable",
                error_details=error_message,
                context={"sessionId": session_id, "state": session_state},
            )
            raise RuntimeError(error_message)

        statement_code = (
            'val _spark = org.apache.spark.sql.SparkSession.builder().getOrCreate(); '
            'val _df = _spark.sql("""' + sql.replace('"""', '\\"\\"\\"') + '"""); '
            '_df.show(20, false); _df'
        )
        statement_response = await client.post(
            f"{LIVY_SERVICE_URL}/sessions/{session_id}/statements",
            json={"code": statement_code},
        )
        if statement_response.status_code >= 300:
            error_message = f"Livy statement submission failed ({statement_response.status_code}): {statement_response.text[:500]}"
            _record_execution_log(
                organization_id=organization_id,
                user_id=user_id,
                job_id=job_id,
                data_source_id=data_source_id,
                execution_engine=execution_engine,
                query_id=f"{job_id or 'spark'}-query",
                stage="spark_statement",
                source="backend",
                status="error",
                message="Livy statement submission failed",
                error_details=error_message,
                context={"sessionId": session_id, "sql": sql},
            )
            raise RuntimeError(error_message)

        statement = statement_response.json()
        _record_execution_log(
            organization_id=organization_id,
            user_id=user_id,
            job_id=job_id,
            data_source_id=data_source_id,
            execution_engine=execution_engine,
            query_id=f"{job_id or 'spark'}-query",
            stage="spark_statement",
            source="backend",
            status="success",
            message="SQL statement submitted to Spark through Livy",
            context={"sessionId": session_id, "statementId": statement.get("id"), "state": statement.get("state")},
        )

        statement_state = (statement.get("state") or "").lower()
        while statement_state not in {"available", "error", "cancelled", "cancelling"}:
            await asyncio.sleep(2)
            statement_response = await client.get(
                f"{LIVY_SERVICE_URL}/sessions/{session_id}/statements/{statement.get('id')}"
            )
            if statement_response.status_code >= 300:
                error_message = f"Livy statement polling failed ({statement_response.status_code}): {statement_response.text[:500]}"
                _record_execution_log(
                    organization_id=organization_id,
                    user_id=user_id,
                    job_id=job_id,
                    data_source_id=data_source_id,
                    execution_engine=execution_engine,
                    query_id=f"{job_id or 'spark'}-query",
                    stage="spark_statement",
                    source="backend",
                    status="error",
                    message="Livy statement polling failed",
                    error_details=error_message,
                    context={"sessionId": session_id, "statementId": statement.get("id")},
                )
                raise RuntimeError(error_message)
            statement = statement_response.json()
            statement_state = (statement.get("state") or "").lower()

        if statement_state in {"error", "cancelled", "cancelling"}:
            output = statement.get("output") or {}
            error_text = output.get("evalue") or output.get("traceback") or json.dumps(output)
            error_message = f"Livy Spark execution failed: {error_text}"
            _record_execution_log(
                organization_id=organization_id,
                user_id=user_id,
                job_id=job_id,
                data_source_id=data_source_id,
                execution_engine=execution_engine,
                query_id=f"{job_id or 'spark'}-query",
                stage="spark_execution",
                source="backend",
                status="error",
                message="Spark execution failed in Livy",
                error_details=error_message,
                context={"sessionId": session_id, "statementId": statement.get("id"), "output": output},
            )
            raise RuntimeError(error_message)

        rows = _extract_rows_from_output(statement.get("output") or {})
        if rows:
            _record_execution_log(
                organization_id=organization_id,
                user_id=user_id,
                job_id=job_id,
                data_source_id=data_source_id,
                execution_engine=execution_engine,
                query_id=f"{job_id or 'spark'}-query",
                stage="spark_execution",
                source="backend",
                status="success",
                message="Spark query completed successfully",
                context={"sessionId": session_id, "rowCount": max(0, len(rows) - 1)},
            )
            return rows

    _record_execution_log(
        organization_id=organization_id,
        user_id=user_id,
        job_id=job_id,
        data_source_id=data_source_id,
        execution_engine=execution_engine,
        query_id=f"{job_id or 'spark'}-query",
        stage="spark_execution",
        source="backend",
        status="success",
        message="Spark query returned no rows but completed",
        context={"sessionId": session_id if 'session_id' in locals() else None},
    )
    return [["Status", "Value"], ["Execution engine", "Spark"], ["Livy session", str(session_id)], ["Connector", connector_type], ["Result", "Executed successfully via Livy"]]


async def _run_report_job(
    job_id: str,
    prompt: str,
    data_source: dict,
    execution_engine: str,
    user: SecurityContext,
):
    data_source = await _resolve_data_source_credentials(data_source, user.tenant_id)
    execution_engine = _normalize_execution_engine(execution_engine, data_source.get("type"))
    org = store.get_org_by_tenant(redis_client, user.tenant_id)
    org_name = org["name"] if org else user.organization_name or user.tenant_id

    steps = [
        (10, "running", "Job queued"),
        (30, "running", "Discovering schema..."),
        (55, "running", "Generating SQL with ADK agent..."),
        (80, "running", "Running safety checks..."),
    ]

    for progress, status, message in steps:
        store.update_job(
            redis_client,
            job_id,
            {"progress": progress, "status": status, "message": message},
        )
        store.upsert_active_job(
            redis_client,
            {
                "id": job_id,
                "organization": org_name,
                "type": "Report generation",
                "progress": progress,
                "status": status,
                "startedAt": "just now",
            },
        )
        await asyncio.sleep(0.7)

    _record_execution_log(
        organization_id=user.tenant_id,
        user_id=user.user_id,
        job_id=job_id,
        data_source_id=data_source.get("id"),
        execution_engine=execution_engine,
        query_id=job_id,
        stage="sql_generated",
        source="backend",
        status="info",
        message="SQL generated for report request",
        context={"prompt": prompt, "dialect": "auto"},
    )
    sql_result = await sql_agent_service.generate_sql(prompt=prompt, data_source=data_source, tenant_id=user.tenant_id)
    if not sql_result.valid:
        failure_message = "SQL validation failed: " + "; ".join(sql_result.validation_errors)
        _record_execution_log(
            organization_id=user.tenant_id,
            user_id=user.user_id,
            job_id=job_id,
            data_source_id=data_source.get("id"),
            execution_engine=execution_engine,
            query_id=job_id,
            stage="sql_validation",
            source="backend",
            status="error",
            message="SQL validation failed",
            error_details=failure_message,
            context={"validationErrors": sql_result.validation_errors, "sql": sql_result.sql},
        )
        store.update_job(
            redis_client,
            job_id,
            {
                "progress": 100,
                "status": "failed",
                "message": failure_message,
                "content": failure_message,
            },
        )
        store.remove_active_job(redis_client, job_id)
        return

    _record_execution_log(
        organization_id=user.tenant_id,
        user_id=user.user_id,
        job_id=job_id,
        data_source_id=data_source.get("id"),
        execution_engine=execution_engine,
        query_id=job_id,
        stage="execution_start",
        source="backend",
        status="info",
        message=f"Starting {execution_engine.upper()} execution",
        context={"sql": sql_result.sql, "dialect": sql_result.dialect},
    )
    try:
        if execution_engine == "spark":
            table_data = await _execute_report_sql_via_livy(
                sql_result.sql,
                data_source,
                organization_id=user.tenant_id,
                user_id=user.user_id,
                job_id=job_id,
                data_source_id=data_source.get("id"),
                execution_engine=execution_engine,
            )
            result_label = f"Executed in Spark via Livy against {data_source['name']}"
        else:
            table_data = _execute_report_sql(sql_result.sql, data_source)
            result_label = f"Executed query against {data_source['name']} on {execution_engine.upper()}"
        row_count = max(0, len(table_data) - 1)
    except Exception as exc:
        error_message = str(exc)
        _record_execution_log(
            organization_id=user.tenant_id,
            user_id=user.user_id,
            job_id=job_id,
            data_source_id=data_source.get("id"),
            execution_engine=execution_engine,
            query_id=job_id,
            stage="execution_failed",
            source="backend",
            status="error",
            message="Report execution failed",
            error_details=error_message,
            context={"sql": sql_result.sql, "dialect": sql_result.dialect},
        )
        table_data = [
            ["Attribute", "Value"],
            ["Data Source", data_source["name"]],
            ["Execution engine", execution_engine.upper()],
            ["Dialect", sql_result.dialect],
            ["Estimated Cost", sql_result.estimated_cost],
            ["PII detected", "yes" if sql_result.pii_detected else "no"],
            ["Provider", sql_result.provider],
            ["Execution error", error_message],
        ]
        result_label = f"SQL generated but execution failed: {exc}"
        row_count = 0

    _record_execution_log(
        organization_id=user.tenant_id,
        user_id=user.user_id,
        job_id=job_id,
        data_source_id=data_source.get("id"),
        execution_engine=execution_engine,
        query_id=job_id,
        stage="execution_complete",
        source="backend",
        status="success",
        message="Report execution completed",
        context={"rowCount": row_count, "resultLabel": result_label},
    )

    content = (
        f'SQL generated for "{prompt}" on {data_source["name"]}.\n\n'
        f'Execution engine: {execution_engine.upper()}\n'
        f'Dialect: {sql_result.dialect}\n'
        f'Cost estimate: {sql_result.estimated_cost}\n'
        f'Validation: passed\n\n'
        f'Generated SQL:\n{sql_result.sql}\n\n'
        f'{result_label}'
    )

    store.update_job(
        redis_client,
        job_id,
        {
            "progress": 100,
            "status": "completed",
            "message": "Report ready",
            "content": content,
            "sqlQuery": sql_result.sql,
            "sqlDialect": sql_result.dialect,
            "sqlValidation": "passed",
            "sqlCost": sql_result.estimated_cost,
            "piiDetected": sql_result.pii_detected,
            "piiColumns": sql_result.pii_columns,
            "plannerNotes": sql_result.planner_notes,
            "provider": sql_result.provider,
            "tableData": table_data,
            "dataSourceName": data_source["name"],
            "prompt": prompt,
            "executionEngine": execution_engine,
        },
    )

    db_user_id = _resolve_db_user_id(user.user_id, email=user.email, username=user.username)
    with SessionLocal() as db:
        job_record = db.query(ReportJob).filter(ReportJob.id == job_id).first()
        if job_record is not None:
            job_record.status = "completed"
            job_record.progress = 100
            job_record.result_summary = {
                "sqlQuery": sql_result.sql,
                "sqlDialect": sql_result.dialect,
                "provider": sql_result.provider,
                "rowCount": row_count,
                "executionEngine": execution_engine,
                "prompt": prompt,
            }
            db.add(job_record)

        report_record = db.query(Report).filter(Report.job_id == job_id).first()
        if report_record is None:
            report_record = Report(
                organization_id=user.tenant_id,
                job_id=job_id,
                title=f"Report: {prompt[:120]}",
                tenant_id=user.tenant_id,
                region=user.region or "unknown",
                content=content,
            )
            db.add(report_record)
        else:
            report_record.title = f"Report: {prompt[:120]}"
            report_record.content = content
            report_record.region = user.region or report_record.region or "unknown"
        db.commit()

    store.remove_active_job(redis_client, job_id)
    store.increment_queries_today(redis_client)

    sources = store.get_data_sources(redis_client, user.tenant_id)
    for source in sources:
        if source["id"] == data_source["id"]:
            source["queryCount"] = source.get("queryCount", 0) + 1
            source["lastUsed"] = datetime.now(timezone.utc).isoformat()
            break
    store.save_data_sources(redis_client, user.tenant_id, sources)

    history_item = {
        "id": job_id,
        "prompt": prompt,
        "dataSourceId": data_source["id"],
        "dataSourceName": data_source["name"],
        "executedAt": datetime.now(timezone.utc).isoformat(),
        "status": "completed",
        "rowCount": row_count,
    }
    store.add_query_history(redis_client, user.tenant_id, history_item)

    notification_id = str(uuid.uuid4())
    store.add_notification(
        redis_client,
        user.tenant_id,
        {
            "id": notification_id,
            "title": f"Report ready: {prompt[:40]}",
            "message": "Your instant report has been generated and is ready to export.",
            "type": "report",
            "read": False,
            "createdAt": datetime.now(timezone.utc).isoformat(),
        },
    )
    with SessionLocal() as db:
        db.add(
            Notification(
                id=notification_id,
                organization_id=user.tenant_id,
                user_id=_resolve_db_user_id(user.user_id, email=user.email, username=user.username),
                title=f"Report ready: {prompt[:40]}",
                message="Your instant report has been generated and is ready to export.",
                type="report",
                is_read=False,
            )
        )
        db.commit()

    orgs = store.get_organizations(redis_client)
    for org in orgs:
        if org["id"] == user.tenant_id:
            org["queriesThisMonth"] = org.get("queriesThisMonth", 0) + 1
            break
    store.save_organizations(redis_client, orgs)


@router.post("/reports/generate")
async def generate_report(
    body: ReportGenerateRequest,
    background_tasks: BackgroundTasks,
    user: SecurityContext = Depends(get_current_user),
):
    tenant_id = user.tenant_id
    sources = store.get_data_sources(redis_client, tenant_id)
    data_source = next((s for s in sources if s["id"] == body.dataSourceId), None)
    if not data_source:
        raise HTTPException(status_code=404, detail="Data source not found")

    data_source = _ensure_data_source_exists(data_source, tenant_id)
    data_source = await _resolve_data_source_credentials(data_source, tenant_id)
    execution_engine = _normalize_execution_engine(body.executionEngine, data_source.get("type"))

    job_id = store.create_job(
        redis_client,
        {
            "status": "running",
            "progress": 0,
            "message": "Starting...",
            "prompt": body.prompt,
            "dataSourceId": body.dataSourceId,
            "tenantId": tenant_id,
            "userId": user.user_id,
            "executionEngine": execution_engine,
        },
    )

    db_user_id = _resolve_db_user_id(user.user_id, email=user.email, username=user.username)
    with SessionLocal() as db:
        db_job = ReportJob(
            id=job_id,
            organization_id=tenant_id,
            user_id=db_user_id or user.user_id,
            data_source_id=data_source.get("id"),
            prompt=body.prompt,
            status="running",
            progress=0,
        )
        db.add(db_job)
        db.commit()

    _record_execution_log(
        organization_id=tenant_id,
        user_id=user.user_id,
        job_id=job_id,
        data_source_id=data_source.get("id"),
        execution_engine=execution_engine,
        query_id=job_id,
        stage="ui_request",
        source="ui",
        status="info",
        message="Report request received from the UI",
        context={"prompt": body.prompt, "dataSourceId": body.dataSourceId, "executionEngine": execution_engine},
        email=user.email,
        username=user.username,
    )

    background_tasks.add_task(_run_report_job, job_id, body.prompt, data_source, execution_engine, user)
    return {"jobId": job_id}


@router.get("/reports/jobs/{job_id}")
async def get_report_job(job_id: str, user: SecurityContext = Depends(get_current_user)):
    job = store.get_job(redis_client, job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    if job.get("tenantId") != user.tenant_id and "PLATFORM_ADMIN" not in user.roles:
        raise HTTPException(status_code=403, detail="Access denied")
    return job


@router.get("/reports/jobs/{job_id}/stream")
async def stream_report_job(job_id: str, user: SecurityContext = Depends(get_current_user)):
    job = store.get_job(redis_client, job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    if job.get("tenantId") != user.tenant_id and "PLATFORM_ADMIN" not in user.roles:
        raise HTTPException(status_code=403, detail="Access denied")

    async def event_stream():
        while True:
            current = store.get_job(redis_client, job_id)
            if not current:
                yield f"event: error\ndata: {json.dumps({'message': 'Job not found'})}\n\n"
                break
            payload = {
                "jobId": job_id,
                "status": current.get("status"),
                "progress": current.get("progress", 0),
                "message": current.get("message", ""),
                "content": current.get("content"),
                "sqlQuery": current.get("sqlQuery"),
                "sqlDialect": current.get("sqlDialect"),
                "sqlValidation": current.get("sqlValidation"),
                "sqlCost": current.get("sqlCost"),
                "piiDetected": current.get("piiDetected"),
                "piiColumns": current.get("piiColumns"),
                "plannerNotes": current.get("plannerNotes"),
                "provider": current.get("provider"),
                "tableData": current.get("tableData"),
                "dataSourceName": current.get("dataSourceName"),
            }
            yield f"data: {json.dumps(payload)}\n\n"
            if current.get("status") in ("completed", "failed"):
                break
            await asyncio.sleep(0.5)

    return StreamingResponse(
        event_stream(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "Connection": "keep-alive"},
    )


@router.get("/reports/jobs/{job_id}/export")
async def export_report_job(
    job_id: str,
    format: str = "csv",
    user: SecurityContext = Depends(get_current_user),
):
    job = store.get_job(redis_client, job_id)
    if not job:
        raise HTTPException(status_code=404, detail="Job not found")
    if job.get("tenantId") != user.tenant_id and "PLATFORM_ADMIN" not in user.roles:
        raise HTTPException(status_code=403, detail="Access denied")
    if job.get("status") != "completed":
        raise HTTPException(status_code=400, detail="Report not ready")

    table: list[list[str]] = job.get("tableData") or []
    filename_base = f"report-{job_id[:8]}"

    if format == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerows(table)
        return StreamingResponse(
            iter([buffer.getvalue()]),
            media_type="text/csv",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.csv"'},
        )

    if format == "excel":
        try:
            from openpyxl import Workbook
        except ImportError as exc:
            raise HTTPException(status_code=501, detail="Excel export unavailable") from exc
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"
        for row in table:
            ws.append(row)
        xlsx_buffer = io.BytesIO()
        wb.save(xlsx_buffer)
        xlsx_buffer.seek(0)
        return StreamingResponse(
            xlsx_buffer,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.xlsx"'},
        )

    if format == "pdf":
        try:
            from fpdf import FPDF
        except ImportError as exc:
            raise HTTPException(status_code=501, detail="PDF export unavailable") from exc
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Helvetica", size=12)
        pdf.cell(0, 10, job.get("content", "Report")[:120], ln=True)
        pdf.ln(4)
        if table:
            col_width = 190 / max(len(table[0]), 1)
            pdf.set_font("Helvetica", size=10)
            for row in table:
                for cell in row:
                    pdf.cell(col_width, 8, str(cell)[:30], border=1)
                pdf.ln()
        pdf_bytes = pdf.output()
        return StreamingResponse(
            io.BytesIO(pdf_bytes),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="{filename_base}.pdf"'},
        )

    raise HTTPException(status_code=400, detail="Unsupported format. Use csv, excel, or pdf")
